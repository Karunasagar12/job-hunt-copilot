"""Generated fictional documents exercise the onboarding upload boundary."""
from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas
from werkzeug.datastructures import FileStorage

import paths


def pdf(text='Jane Doe Senior DS'):
    output = BytesIO()
    page = canvas.Canvas(output)
    if text:
        page.drawString(72, 720, text)
    page.showPage()
    page.save()
    return output.getvalue()


def docx(text='Jane Doe LinkedIn experience', table=False):
    output = BytesIO()
    document = Document()
    document.add_paragraph(text)
    if table:
        document.add_table(rows=1, cols=1).cell(0, 0).text = 'Senior DS table experience'
    document.save(output)
    return output.getvalue()


def upload(client, cv=None, linkedin=None):
    data = {}
    if cv is not None:
        data['cv'] = (BytesIO(cv[0]), cv[1])
    if linkedin is not None:
        data['linkedin'] = (BytesIO(linkedin[0]), linkedin[1])
    return client.post('/api/onboarding/upload', data=data)


def test_pdf_cv_and_docx_linkedin_upload(client):
    raw_cv = pdf()
    response = upload(client, (raw_cv, '../../Jane.PDF'), (docx(), 'linkedin.docx'))
    assert response.status_code == 200
    assert response.json['chars']['cv'] > 10
    assert response.json['chars']['linkedin'] > 5
    assert response.json['warnings'] == []
    import onboarding
    texts = onboarding.uploaded_text()
    assert 'Senior DS' in texts['cv']
    assert texts['linkedin'] == 'Jane Doe LinkedIn experience'
    assert response.json['chars'] == {kind: len(text) for kind, text in texts.items()}
    assert paths.uploads_dir().joinpath('cv.pdf').read_bytes() == raw_cv


@pytest.mark.parametrize('payload,name,expected', [
    (pdf(''), 'cv.pdf', 'image'),
    (b'', 'cv.pdf', 'empty'),
    (b'wrong', 'cv.txt', '.pdf or .docx'),
    (b'broken', 'cv.pdf', 'read'),
    (b'broken', 'cv.docx', 'read'),
    (b'x' * (10 * 1024 * 1024 + 1), 'cv.pdf', '10 MB'),
], ids=['blank', 'empty', 'unsupported', 'broken-pdf', 'broken-docx', 'oversized'])
def test_invalid_cv_returns_readable_400(client, payload, name, expected):
    response = upload(client, (payload, name))
    assert response.status_code == 400
    assert expected in response.json['error']
    assert response.json['status'] == 400
    assert str(paths.data_dir()) not in response.json['error']


@pytest.mark.parametrize('cv', [None, (pdf(), '')])
def test_cv_is_required(client, cv):
    response = upload(client, cv)
    assert response.status_code == 400
    assert 'CV' in response.json['error']


def test_optional_blank_linkedin_warns(client):
    response = upload(client, (pdf(), 'cv.pdf'), (pdf(''), 'linkedin.pdf'))
    assert response.status_code == 200
    assert response.json['chars']['linkedin'] == 0
    assert 'LinkedIn' in response.json['warnings'][0]


@pytest.mark.parametrize('name', ['bad.txt', 'bad.pdf', 'bad.docx'])
def test_invalid_linkedin_does_not_replace_valid_uploads(client, name):
    assert upload(client, (pdf(), 'cv.pdf')).status_code == 200
    before = paths.uploads_dir().joinpath('cv.pdf').read_bytes()
    response = upload(client, (docx(), 'new.docx'), (b'broken', name))
    assert response.status_code == 400
    assert paths.uploads_dir().joinpath('cv.pdf').read_bytes() == before
    assert not paths.uploads_dir().joinpath('cv.docx').exists()


def test_docx_table_and_valid_replacement(client):
    assert upload(client, (pdf(), 'cv.pdf')).status_code == 200
    assert upload(client, (docx(table=True), 'new.DOCX')).status_code == 200
    import onboarding
    assert 'Senior DS table experience' in onboarding.uploaded_text()['cv']
    assert not paths.uploads_dir().joinpath('cv.pdf').exists()


def test_rejected_replacement_preserves_cv(client):
    assert upload(client, (pdf(), 'cv.pdf')).status_code == 200
    assert upload(client, (pdf(''), 'cv.pdf')).status_code == 400
    import onboarding
    assert 'Senior DS' in onboarding.uploaded_text()['cv']


def test_protected_pdf_is_readable_400(client):
    writer = PdfWriter()
    writer.append(PdfReader(BytesIO(pdf())))
    writer.encrypt('fictional-password')
    output = BytesIO()
    writer.write(output)
    response = upload(client, (output.getvalue(), 'cv.pdf'))
    assert response.status_code == 400
    assert 'read' in response.json['error']


def test_helpers_reject_unsupported_type_and_invalid_kind(tmp_path):
    import onboarding
    with pytest.raises(onboarding.OnboardingError, match=r'Upload a .pdf or .docx file.'):
        onboarding.extract_text(tmp_path / 'cv.txt')
    for kind in ('../cv', '/cv', 'other', 'CV'):
        with pytest.raises(onboarding.OnboardingError, match='kind'):
            onboarding.save_upload(FileStorage(stream=BytesIO(pdf()), filename='cv.pdf'), kind)


def test_empty_uploaded_text():
    import onboarding
    assert onboarding.uploaded_text() == {'cv': '', 'linkedin': ''}


def test_combined_request_keeps_existing_json_413(client):
    response = upload(client, (b'x' * (7 * 1024 * 1024), 'cv.pdf'),
                      (b'x' * (7 * 1024 * 1024), 'linkedin.pdf'))
    assert response.status_code == 413
    assert response.json['status'] == 413
