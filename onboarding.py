"""Extract onboarding documents and keep raw uploads in the private data directory."""
from contextlib import contextmanager
from pathlib import Path
from tempfile import NamedTemporaryFile
from copy import deepcopy
import json
import re
from threading import Lock, Thread
from uuid import uuid4

from docx import Document
from pypdf import PdfReader

import paths
import llm_client
import skill_loader

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
EXTENSIONS = ('.pdf', '.docx')
EMPTY_CV_ERROR = 'This PDF looks like an image. Export your CV as a text PDF or DOCX.'


class OnboardingError(ValueError):
    """A document cannot be accepted for onboarding."""


def extract_text(path: Path) -> str:
    extension = path.suffix.lower()
    if extension not in EXTENSIONS:
        raise OnboardingError('Upload a .pdf or .docx file.')
    try:
        if extension == '.pdf':
            reader = PdfReader(path)
            if reader.is_encrypted:
                raise OnboardingError('Could not read this protected PDF. Upload an unprotected PDF or DOCX.')
            return '\n'.join(page.extract_text() or '' for page in reader.pages).strip()
        document = Document(path)
        text = [paragraph.text for paragraph in document.paragraphs]
        text.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        return '\n'.join(text).strip()
    except OnboardingError:
        raise
    except Exception:
        # Parser exceptions may contain private filenames or document contents.
        raise OnboardingError('Could not read this file. Upload a valid text PDF or DOCX.') from None


@contextmanager
def prepared_upload(file_storage, kind):
    """Validate in a temporary file before replacing any previous raw upload."""
    if kind not in ('cv', 'linkedin'):
        raise OnboardingError('Invalid upload kind.')
    extension = Path(file_storage.filename or '').suffix.lower()
    if extension not in EXTENSIONS:
        raise OnboardingError('Upload a .pdf or .docx file.')
    content = file_storage.stream.read(MAX_UPLOAD_BYTES + 1)
    if len(content) > MAX_UPLOAD_BYTES:
        raise OnboardingError('Each upload must be 10 MB or smaller.')
    if not content and kind == 'cv':
        raise OnboardingError('The CV file is empty. Upload a text PDF or DOCX.')
    directory = paths.uploads_dir()
    directory.mkdir(parents=True, exist_ok=True)
    with NamedTemporaryFile(dir=directory, suffix=extension, delete=False) as temporary:
        pending = Path(temporary.name)
        temporary.write(content)
    try:
        text = extract_text(pending) if content else ''
        if kind == 'cv' and not text:
            raise OnboardingError(EMPTY_CV_ERROR)
        yield pending, text
    finally:
        pending.unlink(missing_ok=True)


def store_prepared(pending: Path, kind: str) -> Path:
    """Promote a validated document under a constant name, ignoring its filename."""
    if kind not in ('cv', 'linkedin') or pending.suffix.lower() not in EXTENSIONS:
        raise OnboardingError('Invalid upload kind or file type.')
    destination = paths.uploads_dir() / (kind + pending.suffix.lower())
    pending.replace(destination)
    for extension in EXTENSIONS:
        if extension != destination.suffix:
            destination.with_suffix(extension).unlink(missing_ok=True)
    return destination


def save_upload(file_storage, kind) -> Path:
    with prepared_upload(file_storage, kind) as (pending, _text):
        return store_prepared(pending, kind)


def uploaded_text() -> dict[str, str]:
    result = {'cv': '', 'linkedin': ''}
    for kind in result:
        for extension in EXTENSIONS:
            document = paths.uploads_dir() / (kind + extension)
            if document.is_file():
                result[kind] = extract_text(document) if document.stat().st_size else ''
                break
    return result


def parse_llm_json(text: str) -> dict:
    """Accept an object with optional Markdown fences or surrounding prose."""
    if isinstance(text, str):
        decoder = json.JSONDecoder()
        def matching_brace(start):
            depth = 0
            in_string = escaped = False
            for index in range(start, len(text)):
                char = text[index]
                if in_string:
                    if escaped:
                        escaped = False
                    elif char == '\\':
                        escaped = True
                    elif char == '"':
                        in_string = False
                elif char == '"':
                    in_string = True
                elif char == '{':
                    depth += 1
                elif char == '}':
                    depth -= 1
                    if depth == 0:
                        return index + 1
            return None

        position = 0
        while match := re.search(r'\{', text[position:]):
            start = position + match.start()
            try:
                value, _ = decoder.raw_decode(text[start:])
            except ValueError:
                # Skip an entire malformed brace-delimited region so a nested
                # object cannot be mistaken for a valid top-level response.
                end = matching_brace(start)
                if end is None:
                    break
                position = end
                continue
            if isinstance(value, dict):
                return value
            end = matching_brace(start)
            if end is None:
                break
            position = end
    raise OnboardingError('The profile response did not contain a valid JSON object. Please redraft.')


_ACHIEVEMENT = {'id': '', 'text': '', 'metric': '', 'context': '',
                'needs_metric': False, 'unverified': False}
_ROLE = {**dict.fromkeys(('company', 'title', 'location', 'start', 'end',
                         'description', 'owned'), ''), 'hidden': False,
         'achievements': [_ACHIEVEMENT], 'technologies': [''], 'keywords': ['']}
_SCHEMA = {
    'career': {**dict.fromkeys(('narrative', 'why_looking', 'notice_period',
                              'target_compensation', 'work_authorization'), ''),
               'headline_variants': [{'label': '', 'text': ''}], 'target_roles': ['']},
    'roles': [_ROLE],
    'education': [dict.fromkeys(('degree', 'school', 'dates', 'detail'), '')],
    'projects': [dict.fromkeys(('name', 'context', 'text'), '')],
    'skills': {}, 'certifications': ['']}


def validate_source(data: dict) -> dict:
    """Normalize known fields without coercing malformed nested values."""
    def normalize(value, schema, path):
        if isinstance(schema, dict):
            if not isinstance(value, dict):
                raise OnboardingError(f'Invalid profile field {path}: expected an object.')
            if path == 'skills':
                return {key: normalize(items, [''], 'skills.' + key)
                        for key, items in value.items() if isinstance(key, str)}
            return {key: normalize(value.get(key, default(template)), template,
                                   f'{path}.{key}'.strip('.'))
                    for key, template in schema.items()}
        if isinstance(schema, list):
            if not isinstance(value, list):
                raise OnboardingError(f'Invalid profile field {path}: expected a list.')
            return [normalize(item, schema[0], path) for item in value]
        if type(value) is not type(schema):
            expected = 'boolean' if isinstance(schema, bool) else 'text'
            raise OnboardingError(f'Invalid profile field {path}: expected {expected}.')
        return value

    def default(template):
        return {} if isinstance(template, dict) else [] if isinstance(template, list) else template

    return normalize(data, _SCHEMA, '')


def _company_prefix(company):
    words = re.findall(r'[A-Za-z0-9]+', company)
    if not words:
        return 'XX'
    return (words[0][:2] if len(words) == 1 else words[0][0] + words[1][0]).upper().ljust(2, 'X')


def assign_source_ids(source: dict) -> dict:
    """Assign unique IDs, reserving existing IDs before inserting new achievements."""
    result = deepcopy(source)
    seen = set()
    claimed_prefixes = set()
    role_prefixes = []
    # Existing IDs bind prefixes to their roles, independently of role order.
    for role in result['roles']:
        prefix = None
        for achievement in role['achievements']:
            identifier = achievement.get('id', '')
            match = re.fullmatch(r'([A-Z]{2}\d*)-(\d{2,})', identifier)
            if match and identifier not in seen:
                seen.add(identifier)
                prefix = prefix or match[1]
            else:
                achievement['id'] = ''
        role_prefixes.append(prefix)
        if prefix:
            claimed_prefixes.add(prefix)
    for role, prefix in zip(result['roles'], role_prefixes):
        if prefix is None:
            base = _company_prefix(role['company'])
            prefix = base
            suffix = 2
            while prefix in claimed_prefixes:
                prefix = f'{base}{suffix}'
                suffix += 1
            claimed_prefixes.add(prefix)
        highest = max((int(value.rsplit('-', 1)[1]) for value in seen
                       if value.rsplit('-', 1)[0] == prefix), default=0)
        for achievement in role['achievements']:
            if not achievement.get('id'):
                highest += 1
                identifier = f'{prefix}-{highest:02d}'
                seen.add(identifier)
                achievement['id'] = identifier
    return result


def _numbers(text):
    # Grouping commas are equivalent; decimal points remain significant.
    return set(re.findall(r'\d+\.\d+|\.\d+|\d+', re.sub(r'(?<=\d),(?=\d{3}(?:\D|$))', '', text)))


def flag_unverified(source: dict, evidence: str) -> dict:
    result = deepcopy(source)
    supported = _numbers(evidence)
    for role in result['roles']:
        for achievement in role['achievements']:
            numbers = _numbers(achievement['text'] + ' ' + achievement['metric'])
            achievement['unverified'] = bool(numbers - supported)
    return result


def _call_profile_builder(prompt):
    system = skill_loader.load_skill('profile-builder')
    for attempt in range(2):
        response = llm_client.call_llm(system, prompt, max_tokens=8000, timeout=300)
        try:
            return validate_source(parse_llm_json(response))
        except OnboardingError as exc:
            if attempt:
                raise OnboardingError('The profile response failed JSON/schema validation. Please redraft.') from None
            prompt += '\nValidation error: ' + str(exc) + '\nReturn ONLY JSON matching the schema.'


def draft(cv_text: str, linkedin_text: str = '', existing=None) -> dict:
    prompt = '# CV\n' + cv_text
    if linkedin_text:
        prompt += '\n# LinkedIn\n' + linkedin_text
    known = set()
    if existing is not None:
        existing = validate_source(existing)
        known = {a['id'] for r in existing['roles'] for a in r['achievements'] if a['id']}
        prompt += '\n# Existing profile: keep IDs and user edits\n' + json.dumps(existing)
    result = _call_profile_builder(prompt)
    for role in result['roles']:
        for achievement in role['achievements']:
            if achievement['id'] not in known:
                achievement['id'] = ''
    return flag_unverified(assign_source_ids(result), cv_text + '\n' + linkedin_text)


def validate_answers(answers):
    if not isinstance(answers, dict) or any(not isinstance(key, str) or not isinstance(value, str)
                                            for key, value in answers.items()):
        raise OnboardingError('Gap answers must map achievement IDs to text.')
    return dict(answers)


def apply_gaps(source: dict, answersByID: dict, evidence: str) -> dict:
    original = validate_source(source)
    answers = validate_answers(answersByID)
    known = {a['id'] for r in original['roles'] for a in r['achievements']}
    answers = {key: value for key, value in answers.items() if key in known and value.strip()}
    if not answers:
        return original
    prompt = ('MODE: apply_gaps\n# Current profile\n' + json.dumps(original)
              + '\n# Answers by ID\n' + json.dumps(answers))
    proposed = flag_unverified(_call_profile_builder(prompt),
                               evidence + '\n' + '\n'.join(answers.values()))
    updates = {a['id']: a for r in proposed['roles'] for a in r['achievements'] if a['id'] in answers}
    # The model cannot replace roles, unanswered achievements, context, or user edits.
    for role in original['roles']:
        for achievement in role['achievements']:
            if achievement['id'] in updates:
                update = updates[achievement['id']]
                achievement.update(text=update['text'], metric=update['metric'], needs_metric=False,
                                   unverified=update['unverified'])
    return original


# Deliberately process-local: restarting the server requires a fresh draft.
_tasks = {}
_task_lock = Lock()


def start_task(fn, *args) -> str:
    task_id = str(uuid4())
    with _task_lock:
        _tasks[task_id] = {'status': 'running'}

    def run():
        try:
            status = {'status': 'done', 'result': fn(*args)}
        except llm_client.LLMError as exc:
            status = {'status': 'error', 'error': llm_client._MESSAGES.get(exc.kind, llm_client._MESSAGES['unknown'])}
        except OnboardingError:
            status = {'status': 'error', 'error': 'Profile JSON/schema validation failed. Please redraft.'}
        except Exception:
            status = {'status': 'error', 'error': 'Could not build the profile. Please retry or redraft.'}
        with _task_lock:
            _tasks[task_id] = status

    Thread(target=run, daemon=True).start()
    return task_id


def task_status(task_id: str):
    with _task_lock:
        return deepcopy(_tasks.get(task_id))
